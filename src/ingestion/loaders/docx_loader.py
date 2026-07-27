from pathlib import Path

import docx
from langchain_core.documents import Document

from src.ingestion.loaders.base import BaseLoader
from src.ingestion.exceptions import DocumentLoadError, EmptyDocumentError


class DocxLoader(BaseLoader):
    """
    Word (.docx) loader, via python-docx (new dependency -- no
    already-installed library parses .docx). Joins paragraph text; no
    page concept (docx has no reliable page boundaries in the XML
    itself), same convention as TextLoader/HTMLLoader.
    """

    def load(self, file_path: str) -> list[Document]:

        path = Path(file_path)

        if not path.exists():
            raise DocumentLoadError(f"File not found: {file_path}")

        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise DocumentLoadError(f"'{path.name}' could not be parsed as .docx: {exc}") from exc

        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())

        if not text.strip():
            raise EmptyDocumentError(f"'{path.name}' has no extractable text.")

        return [Document(page_content=text, metadata={"source": path.name})]
