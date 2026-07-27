"""
Tests for the .txt/.html/.docx loaders + their factory/analyzer
registrations (docs/05-roadmap.md v0.2.0 "additional loaders").
"""

import docx
import pytest

from src.ingestion.analysers.factory import AnalyzerFactory
from src.ingestion.analysers.generic_analyzer import GenericAnalyzer
from src.ingestion.exceptions import EmptyDocumentError, UnsupportedDocumentError
from src.ingestion.loaders.docx_loader import DocxLoader
from src.ingestion.loaders.factory import LoaderFactory
from src.ingestion.loaders.html_loader import HTMLLoader
from src.ingestion.loaders.text_loader import TextLoader


def test_loader_factory_resolves_new_extensions(tmp_path):
    assert isinstance(LoaderFactory.build(str(tmp_path / "a.txt")), TextLoader)
    assert isinstance(LoaderFactory.build(str(tmp_path / "a.html")), HTMLLoader)
    assert isinstance(LoaderFactory.build(str(tmp_path / "a.htm")), HTMLLoader)
    assert isinstance(LoaderFactory.build(str(tmp_path / "a.docx")), DocxLoader)


def test_analyzer_factory_resolves_new_extensions(tmp_path):
    for ext in ("txt", "html", "htm", "docx"):
        assert isinstance(AnalyzerFactory.build(str(tmp_path / f"a.{ext}")), GenericAnalyzer)


def test_unsupported_extension_still_rejected(tmp_path):
    with pytest.raises(UnsupportedDocumentError):
        LoaderFactory.build(str(tmp_path / "a.xyz"))


def test_text_loader_reads_real_file(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Hello world.\n\nSecond paragraph.", encoding="utf-8")

    docs = TextLoader().load(str(path))

    assert len(docs) == 1
    assert "Second paragraph" in docs[0].page_content
    assert docs[0].metadata["source"] == "notes.txt"


def test_text_loader_rejects_empty_file(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n  ", encoding="utf-8")

    with pytest.raises(EmptyDocumentError):
        TextLoader().load(str(path))


def test_html_loader_strips_tags_and_scripts(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><style>.x{color:red}</style></head>"
        "<body><script>alert(1)</script><h1>Title</h1><p>Body text.</p></body></html>",
        encoding="utf-8",
    )

    docs = HTMLLoader().load(str(path))

    assert "Title" in docs[0].page_content
    assert "Body text." in docs[0].page_content
    assert "alert" not in docs[0].page_content
    assert "color:red" not in docs[0].page_content


def test_docx_loader_reads_real_docx(tmp_path):
    path = tmp_path / "report.docx"
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(str(path))

    docs = DocxLoader().load(str(path))

    assert "First paragraph." in docs[0].page_content
    assert "Second paragraph." in docs[0].page_content
    assert docs[0].metadata["source"] == "report.docx"


def test_generic_analyzer_classifies_general_with_no_layout_signal(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Some plain text content for analysis.", encoding="utf-8")
    docs = TextLoader().load(str(path))

    analysis = GenericAnalyzer().analyze(docs, str(path))

    assert analysis.document_type == "general"
    assert analysis.layout.has_tables is False
    assert analysis.recommendation.chunk_size > 0
